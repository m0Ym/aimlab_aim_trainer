"""
ReportGenerator 类 - 组员 C 负责
使用 python-docx 自动生成 Word 分析报告
作者学号：请在此处填写你的学号
"""
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from datetime import datetime
from config import REPORTS_DIR


class ReportGenerator:
    """报告生成器类"""
    
    def __init__(self, analyzer, visualizer):
        """
        初始化报告生成器
        
        参数:
            analyzer: DataAnalyzer 实例
            visualizer: DataVisualizer 实例
        """
        self.analyzer = analyzer
        self.visualizer = visualizer
        self.document = None
    
    def create_report(self, player_name='Player', filename=None):
        """
        创建完整的 Word 报告
        
        参数:
            player_name: 玩家姓名
            filename: 文件名（不含扩展名）
            
        返回:
            str: 报告文件完整路径
        """
        if filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"AimLab_Report_{player_name}_{timestamp}"
        
        self.document = Document()
        
        section = self.document.sections[0]
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        
        self._add_title()
        self._add_player_info(player_name)
        self._add_summary()
        self._add_metrics_table()
        self._add_charts()
        self._add_analysis()
        self._add_conclusion()
        
        filepath = f"{REPORTS_DIR}/{filename}.docx"
        self.document.save(filepath)
        
        return filepath
    
    def _add_title(self):
        """添加标题"""
        title = self.document.add_heading('Aim Lab 练枪模拟器\n训练分析报告', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        subtitle = self.document.add_paragraph('Aim Training Performance Report')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_format = subtitle.runs[0].font
        subtitle_format.size = Pt(14)
        subtitle_format.italic = True
        
        self.document.add_paragraph()
    
    def _add_player_info(self, player_name):
        """添加玩家信息"""
        info = self.document.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
        run = info.add_run(f'玩家姓名：{player_name}\n')
        run.font.size = Pt(12)
        run.font.bold = True
        
        run = info.add_run(f'报告生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}\n')
        run.font.size = Pt(12)
        
        run = info.add_run(f'会话 ID: {self.analyzer.data_tracker.session_id}\n')
        run.font.size = Pt(12)
        
        self.document.add_paragraph()
    
    def _add_summary(self):
        """添加概要"""
        self.document.add_heading('1. 训练概要', level=1)
        
        metrics = self.analyzer.get_metrics()
        
        summary = self.document.add_paragraph()
        run = summary.add_run(f'本次训练共射击 {metrics["total_clicks"]} 次，')
        run.font.size = Pt(12)
        
        run = summary.add_run(f'命中 {metrics["hits"]} 次，')
        run.font.size = Pt(12)
        
        run = summary.add_run(f'未命中 {metrics["misses"]} 次，')
        run.font.size = Pt(12)
        
        run = summary.add_run(f'命中率 {metrics["accuracy"]:.1f}%。\n')
        run.font.size = Pt(12)
        run.font.bold = True
        
        run = summary.add_run(f'平均击杀时间 (TTK): {metrics["avg_ttk"]:.1f} 毫秒\n')
        run.font.size = Pt(12)
        
        run = summary.add_run(f'平均偏移距离：{metrics["avg_offset"]:.1f} 像素')
        run.font.size = Pt(12)
        
        self.document.add_paragraph()
    
    def _add_metrics_table(self):
        """添加指标表格"""
        self.document.add_heading('2. 详细数据', level=1)
        
        metrics = self.analyzer.get_metrics()
        
        table = self.document.add_table(rows=7, cols=2)
        table.style = 'Table Grid'
        
        headers = [
            ('总射击次数', str(metrics['total_clicks'])),
            ('命中次数', str(metrics['hits'])),
            ('未命中次数', str(metrics['misses'])),
            ('命中率', f'{metrics["accuracy"]:.1f}%'),
            ('平均 TTK', f'{metrics["avg_ttk"]:.1f} ms'),
            ('平均偏移', f'{metrics["avg_offset"]:.1f} pixels'),
            ('训练时长', f'{self.analyzer.data_tracker.get_session_duration():.1f} s')
        ]
        
        for i, (label, value) in enumerate(headers):
            cell_label = table.cell(i, 0)
            cell_value = table.cell(i, 1)
            
            cell_label.text = label
            cell_value.text = value
            
            for paragraph in cell_label.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
                    run.font.bold = True
            
            for paragraph in cell_value.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
        
        self.document.add_paragraph()
    
    def _add_charts(self):
        """添加图表"""
        self.document.add_heading('3. 可视化分析', level=1)
        
        plots = self.visualizer.get_all_plots()
        
        if 'heatmap' in plots:
            self.document.add_heading('3.1 火力分布热力图', level=2)
            try:
                self.document.add_picture(plots['heatmap'], width=Inches(6))
                caption = self.document.add_paragraph('图 1: 未命中点击分布热力图')
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].font.size = Pt(10)
                caption.runs[0].font.italic = True
                self.document.add_paragraph()
            except Exception as e:
                self.document.add_paragraph(f'图表加载失败：{str(e)}')
        
        if 'reaction_time' in plots:
            self.document.add_heading('3.2 反应时间分析', level=2)
            try:
                self.document.add_picture(plots['reaction_time'], width=Inches(6))
                caption = self.document.add_paragraph('图 2: TTK 变化趋势')
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].font.size = Pt(10)
                caption.runs[0].font.italic = True
                self.document.add_paragraph()
            except Exception as e:
                self.document.add_paragraph(f'图表加载失败：{str(e)}')
        
        if 'accuracy_pie' in plots:
            self.document.add_heading('3.3 命中率统计', level=2)
            try:
                self.document.add_picture(plots['accuracy_pie'], width=Inches(5))
                caption = self.document.add_paragraph('图 3: 命中率饼图')
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].font.size = Pt(10)
                caption.runs[0].font.italic = True
                self.document.add_paragraph()
            except Exception as e:
                self.document.add_paragraph(f'图表加载失败：{str(e)}')
    
    def _add_analysis(self):
        """添加分析建议"""
        self.document.add_heading('4. 分析与建议', level=1)
        
        metrics = self.analyzer.get_metrics()
        
        analysis = self.document.add_paragraph()
        
        if metrics['accuracy'] >= 80:
            run = analysis.add_run('✓ 表现优秀！你的命中率非常高，说明瞄准精度很好。\n')
            run.font.size = Pt(12)
        elif metrics['accuracy'] >= 60:
            run = analysis.add_run('✓ 表现良好！命中率处于中等水平，继续练习可以提高。\n')
            run.font.size = Pt(12)
        else:
            run = analysis.add_run('⚠ 需要改进！命中率较低，建议放慢节奏，提高准确性。\n')
            run.font.size = Pt(12)
        
        if metrics['avg_ttk'] < 300:
            run = analysis.add_run('✓ 反应速度快！平均 TTK 很低，说明反应敏捷。\n')
            run.font.size = Pt(12)
        elif metrics['avg_ttk'] < 500:
            run = analysis.add_run('✓ 反应速度正常！平均 TTK 处于合理范围。\n')
            run.font.size = Pt(12)
        else:
            run = analysis.add_run('⚠ 反应速度有待提高！建议进行专门的反应训练。\n')
            run.font.size = Pt(12)
        
        if metrics['avg_offset'] < 200:
            run = analysis.add_run('✓ 准度控制很好！未命中点距离目标较近。\n')
            run.font.size = Pt(12)
        else:
            run = analysis.add_run('⚠ 准度需要改进！未命中点距离目标较远，建议练习跟枪。\n')
            run.font.size = Pt(12)
    
    def _add_conclusion(self):
        """添加结论"""
        self.document.add_heading('5. 结论', level=1)
        
        conclusion = self.document.add_paragraph()
        run = conclusion.add_run('通过本次训练数据分析，可以看出你的瞄准能力和反应速度。')
        run.font.size = Pt(12)
        
        run = conclusion.add_run('\n\n建议继续进行日常训练，重点关注:')
        run.font.size = Pt(12)
        run.font.bold = True
        
        run = conclusion.add_run('\n• 保持稳定的瞄准节奏')
        run.font.size = Pt(12)
        
        run = conclusion.add_run('\n• 提高对不同距离目标的适应能力')
        run.font.size = Pt(12)
        
        run = conclusion.add_run('\n• 加强移动目标的追踪训练')
        run.font.size = Pt(12)
        
        run = conclusion.add_run('\n\n持续训练将帮助你成为更优秀的 FPS 玩家！')
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.italic = True
